import streamlit as st
import pandas as pd
import json
import os
import qrcode
import io
import base64
import hashlib
import uuid
from PIL import Image
from datetime import datetime

# EVA Inventarisierungs-App – vollstaendige Version mit allen Aenderungen integriert

try:
    import openpyxl
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

try:
    from streamlit_quagga import quagga
    QUAGGA_AVAILABLE = True
except ImportError:
    quagga = None
    QUAGGA_AVAILABLE = False

# Dateien
DATA_FILE = "inventar.json"
ROOMS_FILE = "raeume.json"
CATEGORIES_FILE = "kategorien.json"
LENDING_FILE = "ausleihen.json"
USERS_FILE = "benutzer.json"
WORKSPACES_FILE = "workspaces.json"
IMAGES_DIR = "artikel_bilder"
os.makedirs(IMAGES_DIR, exist_ok=True)

# =============================================
# WORKSPACE HILFSFUNKTIONEN
# =============================================

def load_workspaces():
    return load_json(WORKSPACES_FILE, [])

def save_workspaces(ws):
    save_json(WORKSPACES_FILE, ws)

def get_user_workspaces(benutzername):
    all_ws = load_workspaces()
    return [w for w in all_ws if w["besitzer"] == benutzername or benutzername in w.get("mitglieder", [])]

def get_ws_data_files(ws_id):
    ws_dir = os.path.join("workspaces", ws_id)
    os.makedirs(ws_dir, exist_ok=True)
    return {
        "inventar": os.path.join(ws_dir, "inventar.json"),
        "raeume": os.path.join(ws_dir, "raeume.json"),
        "kategorien": os.path.join(ws_dir, "kategorien.json"),
        "ausleihen": os.path.join(ws_dir, "ausleihen.json"),
        "bilder": os.path.join(ws_dir, "bilder")
    }

def init_workspace_data(ws_id):
    files = get_ws_data_files(ws_id)
    os.makedirs(files["bilder"], exist_ok=True)
    st.session_state.inventar = load_json(files["inventar"], [])
    st.session_state.raeume = load_json(files["raeume"], ["Raum 1"])
    st.session_state.kategorien = load_json(files["kategorien"], DEFAULT_KATEGORIEN)
    st.session_state.ausleihen = load_json(files["ausleihen"], [])
    st.session_state.ws_files = files
    erlaubte = st.session_state.get("erlaubte_raeume", [])
    verfuegbare = [r for r in st.session_state.raeume if not erlaubte or r in erlaubte]
    st.session_state.aktiver_raum = verfuegbare[0] if verfuegbare else None

def save_ws_data(key):
    files = st.session_state.get("ws_files", {})
    if key in files:
        save_json(files[key], st.session_state[key if key != "ausleihen" else "ausleihen"])

def get_ws_image_path(artikel_id):
    bilder_dir = st.session_state.get("ws_files", {}).get("bilder", IMAGES_DIR)
    for ext in ["jpg", "jpeg", "png", "webp"]:
        path = os.path.join(bilder_dir, f"{artikel_id}.{ext}")
        if os.path.exists(path):
            return path
    return None

def save_ws_image(artikel_id, uploaded_file):
    bilder_dir = st.session_state.get("ws_files", {}).get("bilder", IMAGES_DIR)
    os.makedirs(bilder_dir, exist_ok=True)
    ext = uploaded_file.name.split(".")[-1].lower()
    path = os.path.join(bilder_dir, f"{artikel_id}.{ext}")
    with open(path, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return path

def ist_ws_besitzer(ws):
    return ws["besitzer"] == st.session_state.benutzername

def get_ws_rolle(ws):
    if ws["besitzer"] == st.session_state.benutzername:
        return "admin"
    mitglieder = ws.get("mitglieder_rollen", {})
    return mitglieder.get(st.session_state.benutzername, "user")

# Rollen & Rechte
ROLLEN_RECHTE = {
    "admin":   ["inventar", "checkout", "barcode", "statistiken", "export", "benutzerverwaltung"],
    "manager": ["inventar", "checkout", "barcode", "statistiken", "export"],
    "user":    ["inventar", "checkout", "barcode"],
    "viewer":  ["inventar"],
}

ROLLEN_NAMEN = {
    "admin":   "Administrator",
    "manager": "Manager",
    "user":    "Benutzer",
    "viewer":  "Betrachter",
}

def hash_passwort(passwort):
    return hashlib.sha256(passwort.encode()).hexdigest()

def load_users():
    if os.path.exists(USERS_FILE):
        with open(USERS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    default = [
        {"benutzername": "admin", "passwort": hash_passwort("admin123"), "rolle": "admin", "name": "Administrator"}
    ]
    save_json(USERS_FILE, default)
    return default

def hat_recht(recht):
    rolle = st.session_state.get("rolle", "")
    return recht in ROLLEN_RECHTE.get(rolle, [])

def hat_raum_zugriff(raum):
    if st.session_state.get("rolle") == "admin":
        return True
    erlaubte_raeume = st.session_state.get("erlaubte_raeume", [])
    if not erlaubte_raeume:
        return True
    return raum in erlaubte_raeume

def get_erlaubte_raeume():
    if st.session_state.get("rolle") == "admin":
        return st.session_state.raeume
    erlaubte = st.session_state.get("erlaubte_raeume", [])
    if not erlaubte:
        return st.session_state.raeume
    return [r for r in st.session_state.raeume if r in erlaubte]

DEFAULT_KATEGORIEN = ["Elektronik", "Moebel", "Buerobedarf", "Werkzeug", "Sonstiges"]

def load_json(file, default):
    if os.path.exists(file):
        with open(file, "r", encoding="utf-8") as f:
            return json.load(f)
    return default

def save_json(file, data):
    with open(file, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

def save_image(artikel_id, uploaded_file):
    return save_ws_image(artikel_id, uploaded_file)

def get_image_path(artikel_id):
    return get_ws_image_path(artikel_id)

def generate_qr_code(artikel):
    info = f"ID: {artikel['id']}\nName: {artikel['name']}\nRaum: {artikel['raum']}\nMenge: {artikel['menge']}"
    qr = qrcode.QRCode(version=1, box_size=6, border=2)
    qr.add_data(info)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf

# --- Session State Initialisierung ---
if "eingeloggt" not in st.session_state:
    st.session_state.eingeloggt = False
if "benutzername" not in st.session_state:
    st.session_state.benutzername = ""
if "rolle" not in st.session_state:
    st.session_state.rolle = ""
if "benutzer_name" not in st.session_state:
    st.session_state.benutzer_name = ""
if "aktiver_workspace" not in st.session_state:
    st.session_state.aktiver_workspace = None
if "ws_files" not in st.session_state:
    st.session_state.ws_files = {}

# =============================================
# LOGIN
# =============================================
if not st.session_state.eingeloggt:
    st.set_page_config(page_title="Login - Inventarisierungs-App", page_icon="🔐", layout="centered")
    st.title("Inventarisierungs-App")
    st.markdown("Bitte melde dich an oder erstelle einen neuen Account.")
    st.divider()

    login_tab, register_tab = st.tabs(["Anmelden", "Registrieren"])

    with login_tab:
        with st.form("login_form"):
            benutzername = st.text_input("Benutzername")
            passwort = st.text_input("Passwort", type="password")
            login_btn = st.form_submit_button("Anmelden", type="primary", use_container_width=True)

        if login_btn:
            users = load_users()
            user = next((u for u in users if u["benutzername"] == benutzername and u["passwort"] == hash_passwort(passwort)), None)
            if user:
                st.session_state.eingeloggt = True
                st.session_state.benutzername = user["benutzername"]
                st.session_state.rolle = user["rolle"]
                st.session_state.benutzer_name = user["name"]
                st.session_state.erlaubte_raeume = user.get("erlaubte_raeume", [])
                st.rerun()
            else:
                st.error("Benutzername oder Passwort falsch!")

        st.divider()
        st.caption("Noch kein Account? Wechsle zum Tab Registrieren.")

    with register_tab:
        st.markdown("Erstelle deinen persönlichen Account.")
        with st.form("register_form"):
            reg_name = st.text_input("Vollstaendiger Name", placeholder="z.B. Max Mustermann")
            reg_benutzername = st.text_input("Benutzername", placeholder="z.B. max.mustermann")
            reg_passwort = st.text_input("Passwort waehlen", type="password")
            reg_passwort2 = st.text_input("Passwort wiederholen", type="password")
            reg_btn = st.form_submit_button("Account erstellen", type="primary", use_container_width=True)

        if reg_btn:
            users = load_users()
            if not reg_name.strip() or not reg_benutzername.strip() or not reg_passwort.strip():
                st.error("Bitte alle Felder ausfuellen!")
            elif reg_passwort != reg_passwort2:
                st.error("Die Passwoerter stimmen nicht ueberein!")
            elif len(reg_passwort) < 6:
                st.error("Das Passwort muss mindestens 6 Zeichen lang sein!")
            elif any(u["benutzername"] == reg_benutzername.strip() for u in users):
                st.warning("Dieser Benutzername ist bereits vergeben!")
            else:
                neuer_user = {
                    "benutzername": reg_benutzername.strip(),
                    "passwort": hash_passwort(reg_passwort),
                    "rolle": "user",
                    "name": reg_name.strip(),
                    "erlaubte_raeume": [],
                    "registriert_am": datetime.now().strftime("%d.%m.%Y %H:%M")
                }
                users.append(neuer_user)
                save_json(USERS_FILE, users)
                st.success(f"Account '{reg_benutzername}' wurde erstellt! Du kannst dich jetzt anmelden.")
                st.balloons()

        st.divider()
        st.caption("Neue Accounts haben standardmaessig keine Adminrechte.")

    st.stop()

if "edit_artikel_id" not in st.session_state:
    st.session_state.edit_artikel_id = None
if "detail_artikel_id" not in st.session_state:
    st.session_state.detail_artikel_id = None

# =============================================
# WORKSPACE AUSWAHL
# =============================================
if st.session_state.aktiver_workspace is None:
    st.set_page_config(page_title="Workspace - Inventarisierungs-App", page_icon="📂", layout="centered")
    st.title("Workspace auswaehlen")
    st.markdown(f"Willkommen, **{st.session_state.benutzer_name}**! Waehle einen Workspace oder erstelle einen neuen.")
    st.divider()

    user_workspaces = get_user_workspaces(st.session_state.benutzername)

    if user_workspaces:
        st.markdown("### Deine Workspaces")
        for ws in user_workspaces:
            with st.container(border=True):
                col1, col2, col3 = st.columns([3, 1, 1])
                with col1:
                    besitzer_label = "Eigener Workspace" if ws["besitzer"] == st.session_state.benutzername else f"Geteilt von {ws['besitzer']}"
                    st.markdown(f"**{ws['name']}**  \n{besitzer_label} | {len(load_json(get_ws_data_files(ws['id'])['raeume'], []))} Raeume")
                    if ws.get("beschreibung"):
                        st.caption(ws["beschreibung"])
                with col2:
                    if ws.get("passwort"):
                        pw_eingabe = st.text_input("Passwort", type="password", key=f"pw_{ws['id']}")
                    else:
                        pw_eingabe = None
                with col3:
                    if st.button("Oeffnen", key=f"open_{ws['id']}", use_container_width=True):
                        if ws.get("passwort") and hash_passwort(pw_eingabe or "") != ws["passwort"]:
                            st.error("Falsches Workspace-Passwort!")
                        else:
                            st.session_state.aktiver_workspace = ws
                            ws_rolle = get_ws_rolle(ws)
                            st.session_state.rolle = ws_rolle
                            mitglieder_info = ws.get("mitglieder_info", {})
                            user_info = mitglieder_info.get(st.session_state.benutzername, {})
                            st.session_state.erlaubte_raeume = user_info.get("erlaubte_raeume", [])
                            init_workspace_data(ws["id"])
                            st.rerun()

                    if ist_ws_besitzer(ws):
                        if st.button("Loeschen", key=f"del_{ws['id']}", help="Workspace loeschen"):
                            all_ws = load_workspaces()
                            all_ws = [w for w in all_ws if w["id"] != ws["id"]]
                            save_workspaces(all_ws)
                            st.success(f"Workspace '{ws['name']}' geloescht!")
                            st.rerun()
    else:
        st.info("Du hast noch keinen Workspace. Erstelle deinen ersten!")

    st.divider()

    st.markdown("### Neuen Workspace erstellen")
    with st.form("ws_erstellen_form"):
        ws_name = st.text_input("Name", placeholder="z.B. Zuhause, Arbeit, Lager...")
        ws_beschreibung = st.text_input("Beschreibung (optional)")
        ws_passwort = st.text_input("Workspace-Passwort (optional)", type="password",
                                     help="Leer lassen = kein Passwortschutz")
        ws_erstellen_btn = st.form_submit_button("Workspace erstellen", type="primary", use_container_width=True)

    if ws_erstellen_btn:
        if not ws_name.strip():
            st.error("Bitte einen Namen eingeben!")
        else:
            neuer_ws = {
                "id": str(uuid.uuid4()),
                "name": ws_name.strip(),
                "beschreibung": ws_beschreibung.strip(),
                "besitzer": st.session_state.benutzername,
                "passwort": hash_passwort(ws_passwort) if ws_passwort.strip() else None,
                "mitglieder": [],
                "mitglieder_rollen": {},
                "mitglieder_info": {},
                "erstellt_am": datetime.now().strftime("%d.%m.%Y %H:%M")
            }
            all_ws = load_workspaces()
            all_ws.append(neuer_ws)
            save_workspaces(all_ws)
            st.success(f"Workspace '{ws_name}' erstellt!")
            st.rerun()

    st.divider()

    st.markdown("### Workspace beitreten")
    with st.form("ws_beitreten_form"):
        einladungscode = st.text_input("Workspace-ID eingeben", placeholder="Workspace-ID vom Besitzer")
        beitreten_btn = st.form_submit_button("Beitreten", use_container_width=True)

    if beitreten_btn:
        all_ws = load_workspaces()
        gefunden_ws = next((w for w in all_ws if w["id"] == einladungscode.strip()), None)
        if not gefunden_ws:
            st.error("Workspace nicht gefunden!")
        elif st.session_state.benutzername in gefunden_ws.get("mitglieder", []):
            st.warning("Du bist bereits Mitglied dieses Workspaces.")
        elif gefunden_ws["besitzer"] == st.session_state.benutzername:
            st.warning("Das ist dein eigener Workspace.")
        else:
            gefunden_ws["mitglieder"].append(st.session_state.benutzername)
            gefunden_ws.setdefault("mitglieder_rollen", {})[st.session_state.benutzername] = "user"
            save_workspaces(all_ws)
            st.success(f"Du bist dem Workspace '{gefunden_ws['name']}' beigetreten!")
            st.rerun()

    st.divider()
    if st.button("Abmelden", use_container_width=True):
        for key in ["eingeloggt", "benutzername", "rolle", "benutzer_name", "erlaubte_raeume", "aktiver_workspace"]:
            st.session_state[key] = False if key == "eingeloggt" else ([] if key == "erlaubte_raeume" else None if key == "aktiver_workspace" else "")
        st.rerun()

    st.stop()

# Workspace-Daten laden falls noch nicht geschehen
if "inventar" not in st.session_state:
    init_workspace_data(st.session_state.aktiver_workspace["id"])
if "erlaubte_raeume" not in st.session_state:
    st.session_state.erlaubte_raeume = []
if "aktiver_raum" not in st.session_state:
    erlaubte = get_erlaubte_raeume()
    st.session_state.aktiver_raum = erlaubte[0] if erlaubte else None

# --- Seitenlayout ---
st.set_page_config(page_title="Inventarisierungs-App", page_icon="📦", layout="wide")
ws = st.session_state.aktiver_workspace
st.title(f"Inventarisierungs-App | {ws['name']}")
col_title, col_user = st.columns([4, 1])
with col_title:
    st.markdown(f"Workspace von **{ws['besitzer']}** | {ROLLEN_NAMEN.get(st.session_state.rolle, '')}")
with col_user:
    st.markdown(f"**{st.session_state.benutzer_name}**")
    col_a, col_b = st.columns(2)
    with col_a:
        if st.button("Workspaces", use_container_width=True, help="Zurueck zur Workspace-Auswahl"):
            st.session_state.aktiver_workspace = None
            for k in ["inventar", "raeume", "kategorien", "ausleihen", "ws_files", "aktiver_raum", "detail_artikel_id"]:
                if k in st.session_state:
                    del st.session_state[k]
            st.rerun()
    with col_b:
        if st.button("Abmelden", use_container_width=True):
            for key in list(st.session_state.keys()):
                del st.session_state[key]
            st.rerun()

# =============================================
# SIDEBAR
# =============================================
with st.sidebar:

    st.header("Raumverwaltung")
    erlaubte_raeume_sidebar = get_erlaubte_raeume()
    if erlaubte_raeume_sidebar:
        aktiver_raum = st.selectbox(
            "Aktiver Raum",
            erlaubte_raeume_sidebar,
            index=erlaubte_raeume_sidebar.index(st.session_state.aktiver_raum)
            if st.session_state.aktiver_raum in erlaubte_raeume_sidebar else 0
        )
        st.session_state.aktiver_raum = aktiver_raum
    else:
        st.warning("Du hast keinen Zugriff auf Raeume.")

    with st.expander("Neuen Raum hinzufuegen"):
        neuer_raum = st.text_input("Raumname", key="neuer_raum_input")
        if st.button("Raum hinzufuegen"):
            if neuer_raum.strip() == "":
                st.error("Bitte einen Raumnamen eingeben!")
            elif neuer_raum.strip() in st.session_state.raeume:
                st.warning("Dieser Raum existiert bereits!")
            else:
                st.session_state.raeume.append(neuer_raum.strip())
                save_json(st.session_state.ws_files["raeume"], st.session_state.raeume)
                st.session_state.aktiver_raum = neuer_raum.strip()
                st.success(f"Raum '{neuer_raum}' hinzugefuegt!")
                st.rerun()

    with st.expander("Raum umbenennen"):
        if st.session_state.raeume:
            raum_umb = st.selectbox("Raum auswaehlen", st.session_state.raeume, key="raum_umb")
            neuer_raum_n = st.text_input("Neuer Name", value=raum_umb, key="raum_umb_name")
            if st.button("Umbenennen", key="btn_raum_umb"):
                if neuer_raum_n.strip() and neuer_raum_n.strip() != raum_umb:
                    idx = st.session_state.raeume.index(raum_umb)
                    st.session_state.raeume[idx] = neuer_raum_n.strip()
                    for a in st.session_state.inventar:
                        if a["raum"] == raum_umb:
                            a["raum"] = neuer_raum_n.strip()
                    if st.session_state.aktiver_raum == raum_umb:
                        st.session_state.aktiver_raum = neuer_raum_n.strip()
                    save_json(st.session_state.ws_files["raeume"], st.session_state.raeume)
                    save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                    st.success(f"'{raum_umb}' umbenannt zu '{neuer_raum_n.strip()}'!")
                    st.rerun()

    with st.expander("Raum loeschen"):
        if st.session_state.raeume:
            raum_loeschen = st.selectbox("Raum auswaehlen", st.session_state.raeume, key="raum_loeschen")
            if st.button("Raum loeschen", key="btn_raum_loeschen"):
                artikel_im_raum = [a for a in st.session_state.inventar if a["raum"] == raum_loeschen]
                if artikel_im_raum:
                    st.error(f"Raum '{raum_loeschen}' enthaelt noch {len(artikel_im_raum)} Artikel.")
                else:
                    st.session_state.raeume.remove(raum_loeschen)
                    save_json(st.session_state.ws_files["raeume"], st.session_state.raeume)
                    st.session_state.aktiver_raum = st.session_state.raeume[0] if st.session_state.raeume else None
                    st.success(f"Raum '{raum_loeschen}' geloescht!")
                    st.rerun()

    st.divider()

    # --- Kategorieverwaltung ---
    st.header("Kategorien")
    with st.expander("Neue Kategorie"):
        neue_kat = st.text_input("Kategoriename", key="neue_kat_input")
        if st.button("Kategorie hinzufuegen"):
            if neue_kat.strip() == "":
                st.error("Bitte einen Namen eingeben!")
            elif neue_kat.strip() in st.session_state.kategorien:
                st.warning("Kategorie existiert bereits!")
            else:
                st.session_state.kategorien.append(neue_kat.strip())
                save_json(st.session_state.ws_files["kategorien"], st.session_state.kategorien)
                st.success(f"Kategorie '{neue_kat}' hinzugefuegt!")
                st.rerun()

    # AENDERUNG 3: Kategorie umbenennen
    with st.expander("Kategorie umbenennen"):
        if st.session_state.kategorien:
            kat_umb = st.selectbox("Kategorie", st.session_state.kategorien, key="kat_umb")
            neuer_kat_n = st.text_input("Neuer Name", value=kat_umb, key="kat_umb_name")
            if st.button("Umbenennen", key="btn_kat_umb"):
                if neuer_kat_n.strip() and neuer_kat_n.strip() != kat_umb:
                    idx = st.session_state.kategorien.index(kat_umb)
                    st.session_state.kategorien[idx] = neuer_kat_n.strip()
                    for a in st.session_state.inventar:
                        if a["kategorie"] == kat_umb:
                            a["kategorie"] = neuer_kat_n.strip()
                    save_json(st.session_state.ws_files["kategorien"], st.session_state.kategorien)
                    save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                    st.success(f"'{kat_umb}' umbenannt zu '{neuer_kat_n.strip()}'!")
                    st.rerun()

    with st.expander("Kategorie loeschen"):
        if st.session_state.kategorien:
            kat_loeschen = st.selectbox("Kategorie", st.session_state.kategorien, key="kat_loeschen")
            if st.button("Kategorie loeschen", key="btn_kat_loeschen"):
                artikel_mit_kat = [a for a in st.session_state.inventar if a["kategorie"] == kat_loeschen]
                if artikel_mit_kat:
                    st.error(f"Kategorie wird noch von {len(artikel_mit_kat)} Artikel(n) verwendet.")
                else:
                    st.session_state.kategorien.remove(kat_loeschen)
                    save_json(st.session_state.ws_files["kategorien"], st.session_state.kategorien)
                    st.success(f"Kategorie '{kat_loeschen}' geloescht!")
                    st.rerun()

    st.divider()

    # --- Artikel hinzufuegen ---
    st.header("Artikel hinzufuegen")
    st.markdown(f"Raum: **{st.session_state.aktiver_raum}**")
    name = st.text_input("Artikelname")
    kategorie = st.selectbox("Kategorie", st.session_state.kategorien)
    menge = st.number_input("Menge", min_value=1, step=1)
    preis = st.number_input("Preis (EUR)", min_value=0.0, step=0.01, format="%.2f")
    barcode_nr = st.text_input("Barcode-Nr. (optional)", placeholder="z.B. 4012345678901")
    notiz = st.text_area("Notiz (optional)")

    if st.button("Artikel hinzufuegen"):
        if name.strip() == "":
            st.error("Bitte einen Artikelnamen eingeben!")
        elif not st.session_state.aktiver_raum:
            st.error("Bitte zuerst einen Raum anlegen!")
        else:
            neuer_artikel = {
                "id": str(uuid.uuid4()),
                "laufnummer": max((a.get("laufnummer", 0) for a in st.session_state.inventar), default=0) + 1,
                "name": name.strip(),
                "kategorie": kategorie,
                "menge": int(menge),
                "verfuegbar": int(menge),
                "raum": st.session_state.aktiver_raum,
                "preis": round(preis, 2),
                "barcode": barcode_nr.strip(),
                "notiz": notiz.strip(),
                "datum": datetime.now().strftime("%d.%m.%Y %H:%M")
            }
            st.session_state.inventar.append(neuer_artikel)
            save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
            st.success(f"Artikel '{name}' in '{st.session_state.aktiver_raum}' hinzugefuegt!")

# =============================================
# DETAIL-ANSICHT
# =============================================
if st.session_state.detail_artikel_id is not None:
    artikel_detail = next((a for a in st.session_state.inventar if a["id"] == st.session_state.detail_artikel_id), None)
    if artikel_detail:
        with st.container(border=True):
            st.markdown(f"## Artikeldetails: {artikel_detail['name']}")
            col_img, col_info = st.columns([1, 2])

            with col_img:
                img_path = get_image_path(artikel_detail["id"])
                if img_path:
                    st.image(img_path, use_container_width=True)
                else:
                    st.markdown("*Kein Bild vorhanden*")

                uploaded = st.file_uploader("Bild hochladen / aendern", type=["jpg", "jpeg", "png", "webp"],
                                            key=f"upload_{artikel_detail['id']}")
                if uploaded:
                    save_image(artikel_detail["id"], uploaded)
                    st.success("Bild gespeichert!")
                    st.rerun()

            with col_info:
                # AENDERUNG 4: Edit-Modus
                if st.session_state.get("edit_artikel_id") == artikel_detail["id"]:
                    st.markdown("### Artikel bearbeiten")
                    with st.form(key=f"edit_form_{artikel_detail['id']}"):
                        edit_name = st.text_input("Artikelname", value=artikel_detail["name"])
                        edit_kat = st.selectbox("Kategorie", st.session_state.kategorien,
                            index=st.session_state.kategorien.index(artikel_detail["kategorie"])
                            if artikel_detail["kategorie"] in st.session_state.kategorien else 0)
                        col_e1, col_e2 = st.columns(2)
                        with col_e1:
                            edit_menge = st.number_input("Gesamtmenge", min_value=1, step=1, value=int(artikel_detail["menge"]))
                        with col_e2:
                            edit_preis = st.number_input("Preis (EUR)", min_value=0.0, step=0.01, format="%.2f", value=float(artikel_detail["preis"]))
                        edit_raum = st.selectbox("Raum", st.session_state.raeume,
                            index=st.session_state.raeume.index(artikel_detail["raum"])
                            if artikel_detail["raum"] in st.session_state.raeume else 0)
                        edit_barcode = st.text_input("Barcode", value=artikel_detail.get("barcode", "") or "")
                        edit_notiz_f = st.text_area("Notiz", value=artikel_detail.get("notiz", "") or "")
                        col_s, col_ab = st.columns(2)
                        with col_s:
                            speichern = st.form_submit_button("Speichern", type="primary", use_container_width=True)
                        with col_ab:
                            abbrechen = st.form_submit_button("Abbrechen", use_container_width=True)

                    if speichern:
                        if not edit_name.strip():
                            st.error("Artikelname darf nicht leer sein!")
                        else:
                            for a in st.session_state.inventar:
                                if a["id"] == artikel_detail["id"]:
                                    diff = int(edit_menge) - a["menge"]
                                    a["name"] = edit_name.strip()
                                    a["kategorie"] = edit_kat
                                    a["menge"] = int(edit_menge)
                                    a["verfuegbar"] = max(0, a.get("verfuegbar", a["menge"]) + diff)
                                    a["preis"] = round(edit_preis, 2)
                                    a["raum"] = edit_raum
                                    a["barcode"] = edit_barcode.strip()
                                    a["notiz"] = edit_notiz_f.strip()
                                    break
                            save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                            st.session_state.edit_artikel_id = None
                            st.success("Artikel gespeichert!")
                            st.rerun()

                    if abbrechen:
                        st.session_state.edit_artikel_id = None
                        st.rerun()

                else:
                    if st.button("Artikel bearbeiten", key=f"btn_edit_{artikel_detail['id']}"):
                        st.session_state.edit_artikel_id = artikel_detail["id"]
                        st.rerun()

                    verfuegbar = artikel_detail.get("verfuegbar", artikel_detail["menge"])
                    if verfuegbar == artikel_detail["menge"]:
                        status = "Verfuegbar"
                    elif verfuegbar == 0:
                        status = "Ausgeliehen"
                    else:
                        status = f"Teils verfuegbar ({verfuegbar}/{artikel_detail['menge']})"

                    st.markdown(f"""
| Feld | Wert |
|---|---|
| **ID** | {artikel_detail['id']} |
| **Name** | {artikel_detail['name']} |
| **Kategorie** | {artikel_detail['kategorie']} |
| **Raum** | {artikel_detail['raum']} |
| **Gesamtmenge** | {artikel_detail['menge']} |
| **Verfuegbar** | {verfuegbar} |
| **Status** | {status} |
| **Preis** | {artikel_detail['preis']:.2f} EUR |
| **Gesamtwert** | {artikel_detail['preis'] * artikel_detail['menge']:.2f} EUR |
| **Barcode** | {artikel_detail.get('barcode', '-') or '-'} |
| **Hinzugefuegt am** | {artikel_detail['datum']} |
""")

                    st.markdown("**Notiz:**")
                    neue_notiz = st.text_area("Notiz bearbeiten", value=artikel_detail.get("notiz", ""),
                                              key=f"notiz_{artikel_detail['id']}")
                    if st.button("Notiz speichern", key=f"save_notiz_{artikel_detail['id']}"):
                        for a in st.session_state.inventar:
                            if a["id"] == artikel_detail["id"]:
                                a["notiz"] = neue_notiz.strip()
                                break
                        save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                        st.success("Notiz gespeichert!")
                        st.rerun()

            # Ausleihhistorie fuer diesen Artikel
            artikel_ausleihen = [a for a in st.session_state.ausleihen if a["artikel_id"] == artikel_detail["id"]]
            if artikel_ausleihen:
                st.markdown("**Ausleihhistorie dieses Artikels:**")
                df_al = pd.DataFrame(artikel_ausleihen).rename(columns={
                    "person": "Person", "checkout_datum": "Ausgabe",
                    "rueckgabe_erwartet": "Erw. Rueckgabe", "rueckgabe_datum": "Rueckgabe", "status": "Status"
                })
                st.dataframe(df_al[["Person", "Ausgabe", "Erw. Rueckgabe", "Rueckgabe", "Status"]],
                             use_container_width=True, hide_index=True)

            with st.expander("QR-Code anzeigen"):
                qr_buf = generate_qr_code(artikel_detail)
                st.image(qr_buf, width=180)
                st.download_button("QR-Code herunterladen", qr_buf,
                                   f"qr_{artikel_detail['name']}.png", "image/png",
                                   key=f"qr_detail_{artikel_detail['id']}")

            if st.button("Detailansicht schliessen", type="secondary"):
                st.session_state.detail_artikel_id = None
                st.session_state.edit_artikel_id = None
                st.rerun()

    st.divider()

# Tabs dynamisch je nach Rolle
tab_namen = ["Inventar"]
if hat_recht("checkout"): tab_namen.append("Checkout / Check-in")
if hat_recht("barcode"): tab_namen.append("Barcode / QR-Code")
if hat_recht("statistiken"): tab_namen.append("Statistiken")
if hat_recht("export"): tab_namen.append("Export")
if hat_recht("benutzerverwaltung"): tab_namen.append("Benutzerverwaltung")

tabs = st.tabs(tab_namen)
tab_index = {name: i for i, name in enumerate(tab_namen)}

# -----------------------------------------------
# TAB: Inventar
# -----------------------------------------------
with tabs[tab_index["Inventar"]]:
    st.subheader(f"Inventar - {st.session_state.aktiver_raum}")

    col1, col2, col3 = st.columns(3)
    with col1:
        suche = st.text_input("Artikel suchen", placeholder="z.B. Laptop...")
    with col2:
        filter_kategorie = st.selectbox("Kategorie filtern", ["Alle"] + st.session_state.kategorien)
    with col3:
        filter_status = st.selectbox("Status filtern", ["Alle", "Verfuegbar", "Ausgeliehen"])

    raum_inventar = [a for a in st.session_state.inventar if a["raum"] == st.session_state.aktiver_raum and hat_raum_zugriff(a["raum"])]

    if suche:
        raum_inventar = [a for a in raum_inventar if suche.lower() in a["name"].lower()]
    if filter_kategorie != "Alle":
        raum_inventar = [a for a in raum_inventar if a["kategorie"] == filter_kategorie]
    if filter_status == "Verfuegbar":
        raum_inventar = [a for a in raum_inventar if a.get("verfuegbar", a["menge"]) > 0]
    elif filter_status == "Ausgeliehen":
        raum_inventar = [a for a in raum_inventar if a.get("verfuegbar", a["menge"]) < a["menge"]]

    if raum_inventar:
        st.markdown("### Artikel")
        cols_per_row = 3
        for i in range(0, len(raum_inventar), cols_per_row):
            cols = st.columns(cols_per_row)
            for j, a in enumerate(raum_inventar[i:i+cols_per_row]):
                with cols[j]:
                    with st.container(border=True):
                        img_path = get_image_path(a["id"])
                        if img_path:
                            st.image(img_path, use_container_width=True)
                        else:
                            st.markdown("*Kein Bild*")

                        verfuegbar = a.get("verfuegbar", a["menge"])
                        if verfuegbar == a["menge"]:
                            status_icon = "[OK]"
                        elif verfuegbar == 0:
                            status_icon = "[AUS]"
                        else:
                            status_icon = "[TEIL]"

                        st.markdown(f"**{a['name']}** {status_icon}")
                        st.caption(f"{a['kategorie']} | {a['raum']}")
                        st.caption(f"Menge: {a['menge']} | Verfuegbar: {verfuegbar} | {a['preis']:.2f} EUR")

                        # AENDERUNG 5: Detail + Quick-Edit Buttons
                        col_det, col_ed = st.columns(2)
                        with col_det:
                            if st.button("Details", key=f"detail_{a['id']}", use_container_width=True):
                                st.session_state.detail_artikel_id = a["id"]
                                st.session_state.edit_artikel_id = None
                                st.rerun()
                        with col_ed:
                            if st.button("Bearbeiten", key=f"quickedit_{a['id']}", use_container_width=True):
                                st.session_state.detail_artikel_id = a["id"]
                                st.session_state.edit_artikel_id = a["id"]
                                st.rerun()
    else:
        st.info(f"Keine Artikel in Raum '{st.session_state.aktiver_raum}' gefunden.")

    st.divider()
    st.subheader("Artikel loeschen")
    alle_artikel_raum = [a for a in st.session_state.inventar if a["raum"] == st.session_state.aktiver_raum]
    if alle_artikel_raum:
        artikel_namen = [f"ID {a.get('laufnummer', a['id'])} - {a['name']}" for a in alle_artikel_raum]
        ausgewaehlter_artikel = st.selectbox("Artikel auswaehlen", artikel_namen)
        if st.button("Artikel loeschen"):
            # Artikel anhand des Namens identifizieren
            ausgewaehlter_name = ausgewaehlter_artikel.split(" - ", 1)[1]
            zu_loeschen = next((a for a in alle_artikel_raum if a["name"] == ausgewaehlter_name), None)
            if zu_loeschen:
                st.session_state.inventar = [a for a in st.session_state.inventar if a["id"] != zu_loeschen["id"]]
                # AENDERUNG 7: ws_files statt DATA_FILE
                save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                st.success("Artikel wurde geloescht!")
                st.rerun()
    else:
        st.info("Keine Artikel zum Loeschen vorhanden.")

# -----------------------------------------------
# TAB: Checkout / Check-in
# -----------------------------------------------
if hat_recht("checkout"):
    with tabs[tab_index["Checkout / Check-in"]]:
        st.subheader("Checkout & Check-in System")

        co_tab1, co_tab2, co_tab3 = st.tabs(["Checkout (Ausleihe)", "Check-in (Rueckgabe)", "Ausleihhistorie"])

        with co_tab1:
            st.markdown("### Artikel ausleihen")
            verfuegbare_artikel = [a for a in st.session_state.inventar if a.get("verfuegbar", a["menge"]) > 0 and hat_raum_zugriff(a["raum"])]

            if verfuegbare_artikel:
                col1, col2 = st.columns(2)
                with col1:
                    artikel_auswahl = st.selectbox(
                        "Artikel auswaehlen",
                        [f"ID {a.get('laufnummer', a['id'])} - {a['name']} ({a['raum']}) | Verfuegbar: {a.get('verfuegbar', a['menge'])}"
                         for a in verfuegbare_artikel],
                        key="checkout_artikel"
                    )
                    ausgeliehen_idx = [f"ID {a.get('laufnummer', a['id'])} - {a['name']} ({a['raum']}) | Verfuegbar: {a.get('verfuegbar', a['menge'])}"
                                       for a in verfuegbare_artikel].index(artikel_auswahl)
                    artikel_obj = verfuegbare_artikel[ausgeliehen_idx]
                    max_menge = artikel_obj.get("verfuegbar", artikel_obj["menge"])
                    checkout_menge = st.number_input("Menge", min_value=1, max_value=max_menge, step=1, key="checkout_menge")

                with col2:
                    person_name = st.text_input("Name der Person", placeholder="z.B. Max Mustermann")
                    checkout_datum = st.date_input("Ausleihdatum", value=datetime.today())
                    rueckgabe_erwartet = st.date_input("Erwartete Rueckgabe", value=datetime.today())
                    checkout_notiz = st.text_area("Notiz (optional)", key="checkout_notiz")

                if st.button("Auschecken", type="primary"):
                    if not person_name.strip():
                        st.error("Bitte den Namen der Person eingeben!")
                    else:
                        ausleihe = {
                            "ausleihe_id": str(uuid.uuid4()),
                            "artikel_id": artikel_obj["id"],
                            "artikel_name": artikel_obj["name"],
                            "raum": artikel_obj["raum"],
                            "menge": int(checkout_menge),
                            "person": person_name.strip(),
                            "checkout_datum": checkout_datum.strftime("%d.%m.%Y"),
                            "rueckgabe_erwartet": rueckgabe_erwartet.strftime("%d.%m.%Y"),
                            "rueckgabe_datum": None,
                            "status": "ausgeliehen",
                            "notiz": checkout_notiz.strip()
                        }
                        st.session_state.ausleihen.append(ausleihe)
                        save_json(st.session_state.ws_files["ausleihen"], st.session_state.ausleihen)

                        for a in st.session_state.inventar:
                            if a["id"] == artikel_obj["id"]:
                                a["verfuegbar"] = a.get("verfuegbar", a["menge"]) - int(checkout_menge)
                                break
                        save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                        st.success(f"'{artikel_obj['name']}' ({checkout_menge}x) wurde an {person_name} ausgeliehen.")
                        st.rerun()
            else:
                st.info("Keine verfuegbaren Artikel zum Ausleihen vorhanden.")

            st.divider()
            st.markdown("### Aktuell ausgeliehene Artikel")
            aktive_ausleihen = [a for a in st.session_state.ausleihen if a["status"] == "ausgeliehen"]
            if aktive_ausleihen:
                df_aktiv = pd.DataFrame(aktive_ausleihen)
                df_aktiv = df_aktiv.rename(columns={
                    "ausleihe_id": "ID", "artikel_name": "Artikel", "raum": "Raum",
                    "menge": "Menge", "person": "Ausgeliehen an",
                    "checkout_datum": "Ausgabe", "rueckgabe_erwartet": "Rueckgabe erwartet", "notiz": "Notiz"
                })
                st.dataframe(df_aktiv[["ID", "Artikel", "Raum", "Menge", "Ausgeliehen an", "Ausgabe", "Rueckgabe erwartet", "Notiz"]],
                             use_container_width=True, hide_index=True)
            else:
                st.info("Keine aktiven Ausleihen.")

        with co_tab2:
            st.markdown("### Artikel zurueckgeben")
            aktive_ausleihen = [a for a in st.session_state.ausleihen if a["status"] == "ausgeliehen"]

            if aktive_ausleihen:
                ausleihe_auswahl = st.selectbox(
                    "Ausleihe auswaehlen",
                    [f"ID {a['ausleihe_id']} - {a['artikel_name']} ({a['menge']}x) an {a['person']} | Ausgabe: {a['checkout_datum']}"
                     for a in aktive_ausleihen],
                    key="checkin_auswahl"
                )
                ausleihe_idx = [f"ID {a['ausleihe_id']} - {a['artikel_name']} ({a['menge']}x) an {a['person']} | Ausgabe: {a['checkout_datum']}"
                                for a in aktive_ausleihen].index(ausleihe_auswahl)
                ausleihe_obj = aktive_ausleihen[ausleihe_idx]

                col1, col2 = st.columns(2)
                with col1:
                    st.info(f"**Artikel:** {ausleihe_obj['artikel_name']}  \n"
                            f"**Ausgeliehen an:** {ausleihe_obj['person']}  \n"
                            f"**Ausgabe:** {ausleihe_obj['checkout_datum']}  \n"
                            f"**Erwartete Rueckgabe:** {ausleihe_obj['rueckgabe_erwartet']}")
                with col2:
                    rueckgabe_datum = st.date_input("Rueckgabedatum", value=datetime.today(), key="checkin_datum")
                    checkin_notiz = st.text_area("Notiz (optional)", key="checkin_notiz")

                if st.button("Einchecken", type="primary"):
                    for a in st.session_state.ausleihen:
                        if a["ausleihe_id"] == ausleihe_obj["ausleihe_id"]:
                            a["status"] = "zurueckgegeben"
                            a["rueckgabe_datum"] = rueckgabe_datum.strftime("%d.%m.%Y")
                            if checkin_notiz.strip():
                                a["notiz"] = (a.get("notiz", "") or "") + f" | Rueckgabe: {checkin_notiz.strip()}"
                            break
                    save_json(st.session_state.ws_files["ausleihen"], st.session_state.ausleihen)

                    for a in st.session_state.inventar:
                        if a["id"] == ausleihe_obj["artikel_id"]:
                            a["verfuegbar"] = a.get("verfuegbar", 0) + ausleihe_obj["menge"]
                            break
                    save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                    st.success(f"'{ausleihe_obj['artikel_name']}' wurde erfolgreich zurueckgebucht!")
                    st.rerun()
            else:
                st.info("Keine aktiven Ausleihen zur Rueckgabe vorhanden.")

        with co_tab3:
            st.markdown("### Vollstaendige Ausleihhistorie")

            col1, col2 = st.columns(2)
            with col1:
                filter_person = st.text_input("Nach Person suchen", placeholder="z.B. Max Mustermann")
            with col2:
                filter_historie_status = st.selectbox("Status", ["Alle", "ausgeliehen", "zurueckgegeben"])

            historie = st.session_state.ausleihen
            if filter_person:
                historie = [a for a in historie if filter_person.lower() in a["person"].lower()]
            if filter_historie_status != "Alle":
                historie = [a for a in historie if a["status"] == filter_historie_status]

            if historie:
                df_hist = pd.DataFrame(historie)
                df_hist = df_hist.rename(columns={
                    "ausleihe_id": "ID", "artikel_name": "Artikel", "raum": "Raum",
                    "menge": "Menge", "person": "Person", "checkout_datum": "Ausgabe",
                    "rueckgabe_erwartet": "Erw. Rueckgabe", "rueckgabe_datum": "Rueckgabe",
                    "status": "Status", "notiz": "Notiz"
                })
                st.dataframe(df_hist, use_container_width=True, hide_index=True)

                csv_hist = df_hist.to_csv(index=False, sep=";", encoding="utf-8-sig")
                st.download_button("Historie als CSV", csv_hist, "ausleihhistorie.csv", "text/csv")
            else:
                st.info("Keine Eintraege gefunden.")

# -----------------------------------------------
# TAB: Barcode / QR-Code
# -----------------------------------------------
if hat_recht("barcode"):
    with tabs[tab_index["Barcode / QR-Code"]]:
        st.subheader("Barcode & QR-Code")
        bc_tab1, bc_tab2 = st.tabs(["Barcode scannen", "QR-Code generieren"])

        with bc_tab1:
            st.markdown("Scanne einen Barcode mit der Kamera, um einen Artikel zu suchen.")
            try:
                scan_result = quagga()
                if scan_result and scan_result.get("codeResult"):
                    scanned_code = scan_result["codeResult"]["code"]
                    st.success(f"Erkannter Barcode: {scanned_code}")
                    gefundene_artikel = [a for a in st.session_state.inventar if a.get("barcode") == scanned_code]
                    if gefundene_artikel:
                        st.dataframe(pd.DataFrame(gefundene_artikel), use_container_width=True, hide_index=True)
                    else:
                        st.warning("Kein Artikel mit diesem Barcode gefunden.")
            except Exception:
                st.info("Barcode-Scanner nicht verfuegbar. Bitte installiere: pip install streamlit-quagga")
                manueller_code = st.text_input("Barcode manuell eingeben", placeholder="z.B. 4012345678901")
                if manueller_code:
                    gefundene_artikel = [a for a in st.session_state.inventar if a.get("barcode") == manueller_code]
                    if gefundene_artikel:
                        st.success("Artikel gefunden:")
                        st.dataframe(pd.DataFrame(gefundene_artikel), use_container_width=True, hide_index=True)
                    else:
                        st.warning("Kein Artikel gefunden.")

        with bc_tab2:
            st.markdown("Generiere einen QR-Code fuer jeden Artikel zum Ausdrucken.")
            if st.session_state.inventar:
                qr_auswahl = st.selectbox(
                    "Artikel fuer QR-Code auswaehlen",
                    [f"ID {a.get('laufnummer', a['id'])} - {a['name']} ({a['raum']})" for a in st.session_state.inventar]
                )
                qr_idx = [f"ID {a.get('laufnummer', a['id'])} - {a['name']} ({a['raum']})" for a in st.session_state.inventar].index(qr_auswahl)
                artikel_fuer_qr = st.session_state.inventar[qr_idx]
                qr_buf = generate_qr_code(artikel_fuer_qr)
                st.image(qr_buf, caption=f"QR-Code: {artikel_fuer_qr['name']}", width=200)
                # AENDERUNG 6: eindeutiger Key
                st.download_button("QR-Code herunterladen", qr_buf,
                                   f"qr_{artikel_fuer_qr['name']}.png", "image/png",
                                   key=f"qr_tab_{artikel_fuer_qr['id']}")
            else:
                st.info("Noch keine Artikel vorhanden.")

# -----------------------------------------------
# TAB: Statistiken
# -----------------------------------------------
if hat_recht("statistiken"):
    with tabs[tab_index["Statistiken"]]:
        st.subheader("Statistiken")

        gesamt_col, raum_col = st.columns(2)
        with gesamt_col:
            st.markdown("### Gesamtuebersicht")
            if st.session_state.inventar:
                col1, col2, col3, col4 = st.columns(4)
                col1.metric("Artikel", len(st.session_state.inventar))
                col2.metric("Gesamtmenge", sum(a["menge"] for a in st.session_state.inventar))
                col3.metric("Gesamtwert EUR", f"{sum(a['preis'] * a['menge'] for a in st.session_state.inventar):.2f}")
                aktive = len([a for a in st.session_state.ausleihen if a["status"] == "ausgeliehen"])
                col4.metric("Aktive Ausleihen", aktive)

                st.markdown("#### Artikel pro Raum")
                raum_stats = {}
                for a in st.session_state.inventar:
                    raum_stats[a["raum"]] = raum_stats.get(a["raum"], 0) + 1
                st.bar_chart(raum_stats)

                st.markdown("#### Artikel pro Kategorie")
                kat_stats = {}
                for a in st.session_state.inventar:
                    kat_stats[a["kategorie"]] = kat_stats.get(a["kategorie"], 0) + 1
                st.bar_chart(kat_stats)
            else:
                st.info("Noch keine Artikel vorhanden.")

        with raum_col:
            st.markdown(f"### Raum: {st.session_state.aktiver_raum}")
            raum_artikel = [a for a in st.session_state.inventar if a["raum"] == st.session_state.aktiver_raum]
            if raum_artikel:
                col1, col2, col3 = st.columns(3)
                col1.metric("Artikel", len(raum_artikel))
                col2.metric("Menge", sum(a["menge"] for a in raum_artikel))
                col3.metric("Wert EUR", f"{sum(a['preis'] * a['menge'] for a in raum_artikel):.2f}")
            else:
                st.info(f"Keine Artikel in '{st.session_state.aktiver_raum}'.")

# -----------------------------------------------
# TAB: Export / Import
# -----------------------------------------------
if hat_recht("export"):
    with tabs[tab_index["Export"]]:
        st.subheader("Import & Export")

        import_tab, export_tab = st.tabs(["Import", "Export"])

        with import_tab:
            st.markdown("### Artikel aus Excel oder CSV importieren")

            st.markdown("#### 1. Vorlage herunterladen")
            vorlage_df = pd.DataFrame(columns=["name", "kategorie", "menge", "preis", "raum", "barcode", "notiz"])
            col_vl1, col_vl2 = st.columns(2)
            with col_vl1:
                if XLSX_AVAILABLE:
                    vorlage_buf = io.BytesIO()
                    with pd.ExcelWriter(vorlage_buf, engine="openpyxl") as writer:
                        vorlage_df.to_excel(writer, index=False, sheet_name="Inventar")
                    vorlage_buf.seek(0)
                    st.download_button("Excel-Vorlage", data=vorlage_buf,
                        file_name="inventar_vorlage.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True)
            with col_vl2:
                vorlage_csv = vorlage_df.to_csv(index=False, sep=";", encoding="utf-8-sig")
                st.download_button("CSV-Vorlage", data=vorlage_csv,
                    file_name="inventar_vorlage.csv", mime="text/csv", use_container_width=True)
            st.caption("Pflichtfelder: name, menge | Optional: kategorie, preis, raum, barcode, notiz")

            st.divider()
            st.markdown("#### 2. Datei hochladen")
            upload_file = st.file_uploader("Excel (.xlsx) oder CSV auswaehlen",
                type=["xlsx", "csv"], key="import_file")

            import_raum_override = st.selectbox(
                "Ziel-Raum (ueberschreibt 'raum'-Spalte wenn gesetzt)",
                ["--- Spalte aus Datei verwenden ---"] + st.session_state.raeume,
                key="import_raum"
            )
            duplikat_verhalten = st.radio(
                "Bei doppeltem Artikelname",
                ["Ueberspringen", "Trotzdem hinzufuegen"],
                horizontal=True, key="duplikat_verhalten"
            )

            if upload_file is not None:
                try:
                    if upload_file.name.endswith(".xlsx"):
                        df_import = pd.read_excel(upload_file, dtype=str)
                    else:
                        df_import = pd.read_csv(upload_file, sep=None, engine="python", dtype=str)

                    df_import.columns = [c.strip().lower() for c in df_import.columns]
                    df_import = df_import.fillna("")

                    st.markdown("#### 3. Vorschau")
                    st.dataframe(df_import, use_container_width=True, hide_index=True)
                    st.markdown(f"**{len(df_import)} Zeilen** bereit zum Import.")

                    pflicht = ["name", "menge"]
                    fehlende = [f for f in pflicht if f not in df_import.columns]
                    if fehlende:
                        st.error(f"Pflichtfelder fehlen: {', '.join(fehlende)}")
                    else:
                        if st.button("Import starten", type="primary", key="btn_import"):
                            bestehende_namen = [a["name"].lower() for a in st.session_state.inventar]
                            importiert = 0
                            uebersprungen = 0

                            for _, row in df_import.iterrows():
                                artikel_name = str(row.get("name", "")).strip()
                                if not artikel_name:
                                    uebersprungen += 1
                                    continue
                                if duplikat_verhalten == "Ueberspringen" and artikel_name.lower() in bestehende_namen:
                                    uebersprungen += 1
                                    continue

                                if import_raum_override != "--- Spalte aus Datei verwenden ---":
                                    raum_val = import_raum_override
                                else:
                                    raum_val = str(row.get("raum", "")).strip() or st.session_state.aktiver_raum
                                    if raum_val not in st.session_state.raeume:
                                        st.session_state.raeume.append(raum_val)
                                        save_json(st.session_state.ws_files["raeume"], st.session_state.raeume)

                                kat_val = str(row.get("kategorie", "")).strip() or "Sonstiges"
                                if kat_val not in st.session_state.kategorien:
                                    st.session_state.kategorien.append(kat_val)
                                    save_json(st.session_state.ws_files["kategorien"], st.session_state.kategorien)

                                try:
                                    menge_val = max(1, int(float(str(row.get("menge", 1)).replace(",", "."))))
                                except Exception:
                                    menge_val = 1
                                try:
                                    preis_val = round(float(str(row.get("preis", 0)).replace(",", ".")), 2)
                                except Exception:
                                    preis_val = 0.0

                                neuer_artikel = {
                                    "id": str(uuid.uuid4()),
                                    "laufnummer": max((a.get("laufnummer", 0) for a in st.session_state.inventar), default=0) + importiert + 1,
                                    "name": artikel_name,
                                    "kategorie": kat_val,
                                    "menge": menge_val,
                                    "verfuegbar": menge_val,
                                    "raum": raum_val,
                                    "preis": preis_val,
                                    "barcode": str(row.get("barcode", "")).strip(),
                                    "notiz": str(row.get("notiz", "")).strip(),
                                    "datum": datetime.now().strftime("%d.%m.%Y %H:%M")
                                }
                                st.session_state.inventar.append(neuer_artikel)
                                bestehende_namen.append(artikel_name.lower())
                                importiert += 1

                            save_json(st.session_state.ws_files["inventar"], st.session_state.inventar)
                            if importiert > 0:
                                st.success(f"{importiert} Artikel erfolgreich importiert!")
                            if uebersprungen > 0:
                                st.warning(f"{uebersprungen} Zeilen uebersprungen.")
                            st.rerun()

                except Exception as e:
                    st.error(f"Fehler beim Lesen der Datei: {e}")

        with export_tab:
            st.markdown("### Daten exportieren")
            col_exp1, col_exp2 = st.columns(2)
            with col_exp1:
                export_option = st.radio("Was moechtest du exportieren?",
                    ["Aktiver Raum", "Alle Raeume", "Ausleihhistorie"], key="export_option")
            with col_exp2:
                export_format = st.radio("Format",
                    ["CSV", "Excel (.xlsx)"] if XLSX_AVAILABLE else ["CSV"],
                    key="export_format")

            if export_option == "Aktiver Raum":
                export_data = [a for a in st.session_state.inventar if a["raum"] == st.session_state.aktiver_raum]
                filename_base = f"inventar_{st.session_state.aktiver_raum}"
            elif export_option == "Alle Raeume":
                export_data = st.session_state.inventar
                filename_base = "inventar_gesamt"
            else:
                export_data = st.session_state.ausleihen
                filename_base = "ausleihhistorie"

            if export_data:
                df_export = pd.DataFrame(export_data)
                if export_format == "CSV" or not XLSX_AVAILABLE:
                    csv = df_export.to_csv(index=False, sep=";", encoding="utf-8-sig")
                    st.download_button(
                        label=f"CSV herunterladen ({len(export_data)} Eintraege)",
                        data=csv, file_name=f"{filename_base}.csv", mime="text/csv"
                    )
                else:
                    xlsx_buf = io.BytesIO()
                    with pd.ExcelWriter(xlsx_buf, engine="openpyxl") as writer:
                        df_export.to_excel(writer, index=False, sheet_name="Export")
                    xlsx_buf.seek(0)
                    st.download_button(
                        label=f"Excel herunterladen ({len(export_data)} Eintraege)",
                        data=xlsx_buf, file_name=f"{filename_base}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )
            else:
                st.info("Keine Daten zum Exportieren.")

# -----------------------------------------------
# TAB: Benutzerverwaltung (nur Admin)
# -----------------------------------------------
if hat_recht("benutzerverwaltung"):
    with tabs[tab_index["Benutzerverwaltung"]]:
        st.subheader("Benutzerverwaltung")

        ws_aktuell = st.session_state.aktiver_workspace
        all_ws = load_workspaces()
        ws_idx = next((i for i, w in enumerate(all_ws) if w["id"] == ws_aktuell["id"]), None)

        bv_tab1, bv_tab2, bv_tab3 = st.tabs(["Mitglieder", "Workspace teilen", "Workspace-Einstellungen"])

        with bv_tab1:
            st.markdown("### Mitglieder dieses Workspaces")
            users = load_users()
            mitglieder = ws_aktuell.get("mitglieder", [])
            mitglieder_rollen = ws_aktuell.get("mitglieder_rollen", {})
            mitglieder_info = ws_aktuell.get("mitglieder_info", {})

            besitzer_obj = next((u for u in users if u["benutzername"] == ws_aktuell["besitzer"]), {})
            alle_mitglieder = [{
                "Benutzername": ws_aktuell["besitzer"],
                "Name": besitzer_obj.get("name", ws_aktuell["besitzer"]),
                "Rolle": "Besitzer",
                "Erlaubte Raeume": "Alle Raeume"
            }]
            for m in mitglieder:
                m_obj = next((u for u in users if u["benutzername"] == m), {})
                m_info = mitglieder_info.get(m, {})
                alle_mitglieder.append({
                    "Benutzername": m,
                    "Name": m_obj.get("name", m),
                    "Rolle": ROLLEN_NAMEN.get(mitglieder_rollen.get(m, "user"), "user"),
                    "Erlaubte Raeume": ", ".join(m_info.get("erlaubte_raeume", [])) or "Alle Raeume"
                })
            st.dataframe(pd.DataFrame(alle_mitglieder), use_container_width=True, hide_index=True)

            if mitglieder:
                st.divider()
                st.markdown("#### Mitglied entfernen")
                entfernen_auswahl = st.selectbox("Mitglied", mitglieder, key="entfernen_auswahl")
                if st.button("Mitglied entfernen", key="btn_entfernen"):
                    all_ws[ws_idx]["mitglieder"].remove(entfernen_auswahl)
                    all_ws[ws_idx].get("mitglieder_rollen", {}).pop(entfernen_auswahl, None)
                    all_ws[ws_idx].get("mitglieder_info", {}).pop(entfernen_auswahl, None)
                    save_workspaces(all_ws)
                    st.success(f"'{entfernen_auswahl}' wurde entfernt!")
                    st.rerun()

            if mitglieder:
                st.divider()
                st.markdown("#### Raumzugriff & Rolle bearbeiten")
                col1, col2 = st.columns(2)
                with col1:
                    edit_mitglied = st.selectbox("Mitglied", mitglieder, key="edit_mitglied")
                    aktuelle_rolle = mitglieder_rollen.get(edit_mitglied, "user")
                    rolle_optionen = ["manager", "user", "viewer"]
                    neue_rolle = st.selectbox("Rolle", rolle_optionen,
                                              format_func=lambda r: ROLLEN_NAMEN[r],
                                              index=rolle_optionen.index(aktuelle_rolle) if aktuelle_rolle in rolle_optionen else 1,
                                              key="edit_rolle")
                with col2:
                    aktuelle_raeume = mitglieder_info.get(edit_mitglied, {}).get("erlaubte_raeume", [])
                    neue_raeume_edit = st.multiselect("Erlaubte Raeume (leer = alle)",
                                                       st.session_state.raeume,
                                                       default=aktuelle_raeume,
                                                       key="edit_raeume")
                if st.button("Aenderungen speichern", key="btn_edit_mitglied"):
                    all_ws[ws_idx].setdefault("mitglieder_rollen", {})[edit_mitglied] = neue_rolle
                    all_ws[ws_idx].setdefault("mitglieder_info", {}).setdefault(edit_mitglied, {})["erlaubte_raeume"] = neue_raeume_edit
                    save_workspaces(all_ws)
                    st.success("Aenderungen gespeichert!")
                    st.rerun()

        with bv_tab2:
            st.markdown("### Workspace teilen")
            st.markdown("Teile die folgende Workspace-ID mit Personen, die beitreten sollen:")
            st.code(ws_aktuell["id"], language=None)
            st.caption("Die Person kann die ID auf der Workspace-Auswahlseite unter 'Workspace beitreten' eingeben.")

            st.divider()
            st.markdown("#### Benutzer direkt hinzufuegen")
            users = load_users()
            nicht_mitglieder = [u for u in users
                                 if u["benutzername"] != ws_aktuell["besitzer"]
                                 and u["benutzername"] not in ws_aktuell.get("mitglieder", [])]
            if nicht_mitglieder:
                direkt_hinzu = st.selectbox("Benutzer auswaehlen",
                                             [f"{u['name']} ({u['benutzername']})" for u in nicht_mitglieder],
                                             key="direkt_hinzu")
                direkt_idx = [f"{u['name']} ({u['benutzername']})" for u in nicht_mitglieder].index(direkt_hinzu)
                direkt_user = nicht_mitglieder[direkt_idx]
                direkt_rolle = st.selectbox("Rolle zuweisen", ["manager", "user", "viewer"],
                                             format_func=lambda r: ROLLEN_NAMEN[r], key="direkt_rolle")
                direkt_raeume = st.multiselect("Erlaubte Raeume (leer = alle)",
                                                st.session_state.raeume, key="direkt_raeume")
                if st.button("Hinzufuegen", key="btn_direkt_hinzu"):
                    all_ws[ws_idx].setdefault("mitglieder", []).append(direkt_user["benutzername"])
                    all_ws[ws_idx].setdefault("mitglieder_rollen", {})[direkt_user["benutzername"]] = direkt_rolle
                    all_ws[ws_idx].setdefault("mitglieder_info", {})[direkt_user["benutzername"]] = {"erlaubte_raeume": direkt_raeume}
                    save_workspaces(all_ws)
                    st.success(f"'{direkt_user['name']}' wurde hinzugefuegt!")
                    st.rerun()
            else:
                st.info("Alle registrierten Benutzer sind bereits Mitglied.")

        with bv_tab3:
            st.markdown("### Workspace-Einstellungen")
            neuer_ws_name = st.text_input("Workspace-Name", value=ws_aktuell["name"])
            neue_ws_beschreibung = st.text_input("Beschreibung", value=ws_aktuell.get("beschreibung", ""))
            neues_ws_passwort = st.text_input("Neues Passwort (leer lassen = unveraendert)", type="password")
            pw_entfernen = st.checkbox("Passwortschutz entfernen")

            if st.button("Einstellungen speichern"):
                all_ws[ws_idx]["name"] = neuer_ws_name.strip()
                all_ws[ws_idx]["beschreibung"] = neue_ws_beschreibung.strip()
                if pw_entfernen:
                    all_ws[ws_idx]["passwort"] = None
                elif neues_ws_passwort.strip():
                    all_ws[ws_idx]["passwort"] = hash_passwort(neues_ws_passwort.strip())
                save_workspaces(all_ws)
                st.session_state.aktiver_workspace = all_ws[ws_idx]
                st.success("Einstellungen gespeichert!")
                st.rerun()

            st.divider()
            st.markdown("#### Workspace loeschen")
            st.warning("Diese Aktion kann nicht rueckgaengig gemacht werden!")
            if st.button("Workspace endgueltig loeschen", type="primary"):
                all_ws = [w for w in all_ws if w["id"] != ws_aktuell["id"]]
                save_workspaces(all_ws)
                st.session_state.aktiver_workspace = None
                for k in ["inventar", "raeume", "kategorien", "ausleihen", "ws_files", "aktiver_raum"]:
                    if k in st.session_state:
                        del st.session_state[k]
                st.rerun()

st.markdown("---")
st.markdown("EVA - Inventarisierungs-App | Erstellt mit Python & Streamlit")
